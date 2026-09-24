from __future__ import annotations

import gc
import json
import re
import sys
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

import httpx
import pandas as pd
import pyarrow.parquet as pq
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from ats_scrapers import Client

# ============================================================
# Configuration
# ============================================================

DAYS_TO_KEEP = 7
OUTPUT_FILE = Path("jobs_latest.xlsx")
JSON_OUTPUT = Path("docs/jobs_latest.json")
HTTP_TIMEOUT = 90.0  # seconds per ATS download

PREFERRED_ATS = [
    "greenhouse",
    "lever",
    "ashby",
    "workable",
    "smartrecruiters",
    "personio",
    "ycombinator",
    "wellfound",
    "weworkremotely",
    "gem",
    "jobbankca",
    "welcometothejungle",
    "join_com",
    "softgarden",
    "teamtailor",
    "recruitee",
    "breezy",
    "bamboohr",
    "pinpoint",
    "eightfold",
    "jazzhr",
    "phenom",
    "icims",
    "rippling",
    "dayforce",
    "jobvite",
    "avature",
    "google",
    "apple",
    "amazon",
    "tesla",
    "uber",
    "builtin",
    "mercor",
]

JOB_KEYWORDS = [
    "application security engineer",
    "senior application security engineer",
    "application security",
    "product security engineer",
    "senior product security engineer",
    "product security",
    "security engineer",
    "senior security engineer",
    "software security engineer",
    "cloud security engineer",
    "cloud application security",
    "devsecops engineer",
    "devsecops security engineer",
    "security software engineer",
    "application security architect",
    "product security architect",
    "cloud security architect",
    "api security engineer",
    "security automation engineer",
    "platform security engineer",
    "software security",
]

TITLE_INCLUDE_TERMS = [
    "application security",
    "product security",
    "security engineer",
    "software security",
    "cloud security",
    "devsecops",
    "security architect",
    "api security",
    "security automation",
    "platform security",
    "appsec",
    "prodsec",
    "threat modeling",
    "threat modelling",
    "vulnerability management",
    "security research",
    "infrastructure security",
    "security platform",
]

EXCLUDE_TITLE_TERMS = [
    "soc",
    "security operations",
    "grc",
    "governance",
    "risk and compliance",
    "compliance",
    "it support",
    "iam administrator",
    "identity administrator",
    "network security",
    "physical security",
    "security guard",
    "security officer",
    "cybersecurity analyst",
    "security analyst",
]

TARGET_COUNTRIES = [
    "austria", "belgium", "bulgaria", "croatia", "cyprus",
    "czech republic", "czechia", "denmark", "estonia", "finland",
    "france", "germany", "greece", "hungary", "iceland",
    "ireland", "italy", "latvia", "lithuania", "luxembourg",
    "malta", "netherlands", "norway", "poland", "portugal",
    "romania", "slovakia", "slovenia", "spain", "sweden",
    "switzerland", "united kingdom", "uk", "england", "scotland",
    "wales", "northern ireland", "canada",
]

TARGET_COUNTRY_ISO = {
    "AT", "BE", "BG", "HR", "CY", "CZ", "DK", "EE", "FI", "FR",
    "DE", "GR", "HU", "IS", "IE", "IT", "LV", "LT", "LU", "MT",
    "NL", "NO", "PL", "PT", "RO", "SK", "SI", "ES", "SE", "CH",
    "GB", "UK", "CA",
}

REMOTE_TERMS = [
    "remote", "fully remote", "work from home", "remote-first",
    "remote work", "work remotely",
]

VISA_POSITIVE_TERMS = [
    "visa sponsorship", "visa sponsor", "visa support",
    "sponsorship available", "sponsorship provided", "visa assistance",
    "immigration support", "immigration sponsorship", "work visa",
    "work permit", "relocation assistance", "relocation support",
    "relocation package", "skilled worker", "skilled worker visa",
    "eu blue card", "highly skilled migrant", "lmia",
    "employer sponsorship", "will sponsor", "sponsors visas",
]

VISA_NEGATIVE_TERMS = [
    "no visa sponsorship", "visa sponsorship is not available",
    "we do not sponsor", "we don't sponsor", "unable to sponsor",
    "cannot sponsor", "must have the right to work",
    "must already have the right to work", "without sponsorship",
    "no sponsorship", "does not sponsor", "not able to sponsor",
]

COLUMNS_NEEDED = [
    "title", "company", "location", "country_iso", "region",
    "is_remote", "salary_summary", "salary_min", "salary_max",
    "salary_currency", "employment_type", "description",
    "posted_at", "apply_url", "url", "ats_type",
]


@dataclass
class Job:
    posted_date: str
    company: str
    title: str
    location: str
    country: str
    remote: str
    visa: str
    salary: str
    experience: str
    ats: str
    apply_url: str
    job_url: str
    description: str


def log(msg: str) -> None:
    print(msg, flush=True)


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return " ".join(text(item) for item in value)
    if isinstance(value, dict):
        return " ".join(text(v) for v in value.values())
    if isinstance(value, float) and pd.isna(value):
        return ""
    return str(value)


def first_value(item: dict[str, Any], *keys: str) -> str:
    for key in keys:
        value = text(item.get(key)).strip()
        if value and value.lower() not in ("nan", "none", "null"):
            return value
    return ""


def parse_date(value: Any) -> datetime | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, datetime):
        if value.tzinfo is None:
            return value.replace(tzinfo=timezone.utc)
        return value.astimezone(timezone.utc)
    raw = str(value).strip()
    if not raw or raw.lower() in ("nan", "none", "null"):
        return None
    raw = raw.replace("Z", "+00:00")
    try:
        parsed = datetime.fromisoformat(raw)
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.astimezone(timezone.utc)
    except ValueError:
        pass
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(raw[:19], fmt).replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    return None


def _title_matches(title: str) -> bool:
    t = title.lower()
    if any(ex in t for ex in EXCLUDE_TITLE_TERMS):
        return False
    return any(inc in t for inc in TITLE_INCLUDE_TERMS)


def _location_ok(row: dict[str, Any]) -> bool:
    loc = first_value(row, "location", "region").lower()
    iso = first_value(row, "country_iso").upper().strip()

    non_target_names = [
        "united states", "usa", "u.s.", "u.s.a", "brazil", "são paulo", "sao paulo",
        "india", "bangalore", "bengaluru", "hyderabad", "china", "japan",
        "australia", "sydney", "melbourne", "singapore", "mexico",
        "argentina", "americas", "latin america", "south america",
        "san francisco", "new york", "seattle", "austin", "boston",
        "los angeles", "chicago", "denver", "atlanta", "miami",
        "us remote", "remote us", "remote (us", "remote - us",
    ]
    us_state_patterns = [
        ", ca", " ca,", "california", ", ny", " ny,", ", wa", " wa,",
        ", tx", " tx,", ", ma", ", co", ", il", ", fl", ", ga",
        ", nj", ", va", ", nc", ", or", ", ut",
    ]

    if any(n in loc for n in non_target_names):
        return False
    if any(p in loc for p in us_state_patterns):
        return False

    non_target_iso = {
        "US", "USA", "BR", "IN", "CN", "JP", "AU", "NZ", "MX", "AR",
        "CL", "CO", "SG", "HK", "TW", "KR", "PH", "ID", "MY", "TH",
        "AE", "SA", "IL", "ZA", "NG", "KE", "EG", "RU", "UA",
    }
    if iso and iso in non_target_iso:
        return False

    if iso and iso in TARGET_COUNTRY_ISO:
        return True

    if any(c in loc for c in TARGET_COUNTRIES):
        return True

    if any(term in loc for term in REMOTE_TERMS + ["anywhere", "worldwide", "global", "europe"]):
        return True

    if row.get("is_remote") is True and (not loc or loc in ("remote", "fully remote", "remote-first")):
        return True

    return False


def _is_recent(posted: Any, cutoff: datetime) -> bool:
    dt = parse_date(posted)
    if dt is None:
        return False
    return dt >= cutoff


def fetch_ats_parquet(url: str) -> pd.DataFrame | None:
    try:
        with httpx.Client(timeout=HTTP_TIMEOUT, follow_redirects=True) as client:
            resp = client.get(url)
            resp.raise_for_status()
        buf = BytesIO(resp.content)
        pf = pq.ParquetFile(buf)
        available = set(pf.schema_arrow.names)
        cols = [c for c in COLUMNS_NEEDED if c in available]
        if "title" not in cols:
            return None
        table = pf.read(columns=cols)
        return table.to_pandas()
    except Exception as exc:
        log(f"  Download/parse failed: {exc}")
        return None


def search_jobs() -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    now = datetime.now(timezone.utc)
    cutoff = now - timedelta(days=DAYS_TO_KEEP)

    log("Starting job search (ATS-by-ATS)...")
    client = Client()
    manifest = client.manifest
    by_ats = manifest.by_ats

    ordered = []
    seen = set()
    for name in PREFERRED_ATS:
        if name in by_ats and by_ats[name].rows > 0:
            ordered.append(name)
            seen.add(name)

    total = len(ordered)
    log(f"Will process {total} ATS sources.")

    for idx, ats_name in enumerate(ordered, 1):
        entry = by_ats[ats_name]
        log(f"[{idx}/{total}] Searching ATS: {ats_name} ({entry.rows} jobs)")

        try:
            df = fetch_ats_parquet(entry.parquet)
            if df is None or df.empty:
                log(f"  Found: 0 (empty or failed)")
                continue

            if "posted_at" in df.columns:
                mask_recent = df["posted_at"].apply(
                    lambda x: _is_recent(x, cutoff)
                )
                df = df[mask_recent]
            if df.empty:
                log(f"  Found: 0 (no recent jobs)")
                del df
                gc.collect()
                continue

            if "title" in df.columns:
                mask_title = df["title"].fillna("").astype(str).apply(_title_matches)
                df = df[mask_title]
            if df.empty:
                log(f"  Found: 0 (no matching titles)")
                del df
                gc.collect()
                continue

            records = df.to_dict(orient="records")
            matched = []
            for rec in records:
                if _location_ok(rec):
                    if not rec.get("ats_type"):
                        rec["ats_type"] = ats_name
                    matched.append(rec)

            log(f"  Found: {len(matched)}")
            results.extend(matched)

            del df, records, matched
            gc.collect()

        except Exception as exc:
            log(f"  Failed for {ats_name}: {exc}")
            continue

    log(f"Raw results collected: {len(results)}")
    return results


def filter_jobs(
    raw_jobs: list[dict[str, Any]],
    cutoff: datetime,
) -> list[dict[str, Any]]:
    filtered = []
    for job in raw_jobs:
        title = first_value(job, "title", "job_title", "name")
        if not title:
            continue
        if not _is_recent(
            first_value(job, "posted_at", "date_posted", "posted_date", "created_at"),
            cutoff,
        ):
            continue
        if not _location_ok(job):
            continue
        if not _title_matches(title):
            continue
        filtered.append(job)
    log(f"After filtering: {len(filtered)}")
    return filtered


def detect_visa_status(job: dict[str, Any]) -> str:
    description = first_value(job, "description", "job_description", "content").lower()
    location = first_value(job, "location", "country").lower()
    combined = f"{description} {location}"

    if any(term in combined for term in VISA_NEGATIVE_TERMS):
        return "No"
    if any(term in combined for term in VISA_POSITIVE_TERMS):
        return "Yes / Mentioned"
    return "Not mentioned"


def normalize_job(job: dict[str, Any]) -> Job:
    title = first_value(job, "title", "job_title", "name")
    company = first_value(job, "company", "company_name", "employer")
    location = first_value(job, "location", "locations", "region")
    country = first_value(job, "country_iso", "country")
    description = first_value(job, "description", "job_description", "content")
    posted = first_value(job, "posted_at", "date_posted", "posted_date", "created_at")
    apply_url = first_value(job, "apply_url", "application_url")
    job_url = first_value(job, "url", "job_url", "link")
    if not apply_url:
        apply_url = job_url

    remote_flag = job.get("is_remote")
    if remote_flag is True:
        remote = "Yes"
    elif any(term in f"{location} {description}".lower() for term in REMOTE_TERMS):
        remote = "Yes"
    else:
        remote = "No"

    salary = first_value(job, "salary_summary", "salary", "salary_text")
    if not salary:
        smin = job.get("salary_min")
        smax = job.get("salary_max")
        curr = first_value(job, "salary_currency")
        if smin or smax:
            parts = []
            if smin and not (isinstance(smin, float) and pd.isna(smin)):
                parts.append(str(int(smin)))
            if smax and not (isinstance(smax, float) and pd.isna(smax)):
                parts.append(str(int(smax)))
            salary = " - ".join(parts)
            if curr:
                salary = f"{salary} {curr}".strip()

    return Job(
        posted_date=posted,
        company=company or "Unknown",
        title=title,
        location=location or "Not specified",
        country=country or "Not specified",
        remote=remote,
        visa=detect_visa_status(job),
        salary=salary,
        experience=first_value(job, "experience", "experience_level", "employment_type"),
        ats=first_value(job, "ats_type", "ats", "source") or "unknown",
        apply_url=apply_url,
        job_url=job_url,
        description=description[:15000] if description else "",
    )


def deduplicate_jobs(jobs: list[Job]) -> list[Job]:
    unique: dict[tuple[str, str, str], Job] = {}
    for job in jobs:
        key = (
            job.company.strip().lower(),
            job.title.strip().lower(),
            job.location.strip().lower(),
        )
        if key not in unique:
            unique[key] = job
    result = list(unique.values())
    log(f"After deduplication: {len(result)}")
    return result


def load_existing_statuses() -> dict[tuple[str, str, str], str]:
    statuses: dict[tuple[str, str, str], str] = {}
    if not OUTPUT_FILE.exists():
        log("No existing workbook found. All jobs will start as New.")
        return statuses
    try:
        workbook = load_workbook(OUTPUT_FILE, read_only=True, data_only=True)
        if "Latest Jobs" not in workbook.sheetnames:
            workbook.close()
            return statuses
        sheet = workbook["Latest Jobs"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        company_col = headers.get("Company")
        title_col = headers.get("Job Title")
        location_col = headers.get("Location")
        status_col = headers.get("Status")
        if not all([company_col, title_col, location_col, status_col]):
            workbook.close()
            return statuses
        for row in sheet.iter_rows(min_row=2, values_only=False):
            company = text(row[company_col - 1].value).strip()
            title = text(row[title_col - 1].value).strip()
            location = text(row[location_col - 1].value).strip()
            status = text(row[status_col - 1].value).strip()
            if company and title and status:
                statuses[(company.lower(), title.lower(), location.lower())] = status
        workbook.close()
        log(f"Loaded {len(statuses)} existing statuses.")
    except Exception as e:
        log(f"Could not load existing statuses: {e}")
    return statuses


def create_workbook(jobs: list[Job], existing_statuses: dict[tuple[str, str, str], str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Latest Jobs"

    headers = [
        "Posted Date", "Company", "Job Title", "Location", "Country",
        "Remote", "Visa/Sponsorship", "Salary", "Experience", "ATS",
        "Apply", "Status", "Job Description",
    ]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    for col in range(1, len(headers) + 1):
        cell = ws.cell(1, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for job in jobs:
        key = (job.company.lower(), job.title.lower(), job.location.lower())
        status = existing_statuses.get(key, "New")
        posted = job.posted_date
        dt = parse_date(posted)
        if dt:
            posted = dt.strftime("%Y-%m-%d")
        ws.append([
            posted,
            job.company,
            job.title,
            job.location,
            job.country,
            job.remote,
            job.visa,
            job.salary,
            job.experience,
            job.ats,
            job.apply_url or job.job_url,
            status,
            job.description,
        ])

    # Apply column is hyperlink
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row, 11)
        url = text(cell.value)
        if url.startswith("http"):
            cell.hyperlink = url
            cell.font = Font(color="0563C1", underline="single")

    # Status dropdown
    dv = DataValidation(type="list", formula1='"New,Applied,Interview,Rejected,Offer,Skip"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"L2:L{max(ws.max_row, 2)}")

    widths = [12, 22, 36, 22, 10, 10, 16, 16, 14, 12, 40, 12, 50]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    wb.save(OUTPUT_FILE)
    log(f"Created: {OUTPUT_FILE} ({len(jobs)} jobs)")


def write_json(jobs: list[Job]) -> None:
    """Write lightweight JSON for the GitHub Pages site (no descriptions)."""
    JSON_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    payload = []
    for job in jobs:
        posted = job.posted_date
        dt = parse_date(posted)
        if dt:
            posted = dt.strftime("%Y-%m-%d")
        payload.append({
            "title": job.title,
            "company": job.company,
            "location": job.location,
            "country": job.country,
            "remote": job.remote,
            "visa": job.visa,
            "posted_date": posted,
            "apply_url": job.apply_url or job.job_url,
        })
    with open(JSON_OUTPUT, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    log(f"Created: {JSON_OUTPUT} ({len(payload)} jobs)")


def main() -> None:
    now = datetime.now(timezone.utc)
    cutoff = now - timedelta(days=DAYS_TO_KEEP)

    raw_jobs = search_jobs()
    filtered_jobs = filter_jobs(raw_jobs, cutoff)
    normalized_jobs = [normalize_job(job) for job in filtered_jobs]
    jobs = deduplicate_jobs(normalized_jobs)

    jobs.sort(
        key=lambda job: parse_date(job.posted_date) or datetime.min.replace(tzinfo=timezone.utc),
        reverse=True,
    )

    existing_statuses = load_existing_statuses()
    create_workbook(jobs, existing_statuses)
    write_json(jobs)
    log("Job search completed successfully.")


if __name__ == "__main__":
    main()
